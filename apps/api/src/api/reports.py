"""Report assembly and export service."""

import io
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader
from weasyprint import HTML

from api.charting import ChartBuilder
from api.deepseek import generate_report_narrative
from api.models.analysis import AnalysisPack, ProxyResult
from api.models.analysis_tables import AnalysisResult
from api.models.interpretation_tables import InterpretationResult

_TEMPLATES_DIR = Path(__file__).resolve().parent / "templates"
_env = Environment(loader=FileSystemLoader(str(_TEMPLATES_DIR)))


def _proxy_label(pack: AnalysisPack, proxy_id: str) -> str:
    for p in pack.proxies:
        if p.id == proxy_id:
            return p.label
    return proxy_id


def _proxy_description(pack: AnalysisPack, proxy_id: str) -> str:
    for p in pack.proxies:
        if p.id == proxy_id:
            return p.description or ""
    return ""


class ReportAssembler:
    """Assembles dataset metadata, proxy table, charts, interpretation into export formats."""

    @staticmethod
    def build_report_data(
        pack: AnalysisPack,
        analysis_rows: list[AnalysisResult],
        interpretation: InterpretationResult | None,
        deepseek_data: dict | None = None,
    ) -> dict[str, Any]:
        """Build the structured data dict for templates."""
        proxy_results = [
            ProxyResult(
                proxy_id=r.proxy_id,
                label=r.proxy_label,
                sample_id=r.sample_id,
                value=r.value,
                unit=r.unit,
            )
            for r in analysis_rows
        ]

        # Build proxy table
        sample_ids = sorted({r.sample_id for r in proxy_results})
        proxy_ids = sorted({r.proxy_id for r in proxy_results})
        proxy_labels = [_proxy_label(pack, pid) for pid in proxy_ids]

        table_rows: list[dict[str, Any]] = []
        for sid in sample_ids:
            row: dict[str, Any] = {"sample_id": sid}
            for pid in proxy_ids:
                matching = [r for r in proxy_results if r.sample_id == sid and r.proxy_id == pid]
                if matching:
                    row[pid] = matching[0].value
                    row[f"{pid}_unit"] = matching[0].unit or ""
                else:
                    row[pid] = None
                    row[f"{pid}_unit"] = ""
            table_rows.append(row)

        # Build charts as Plotly JSON (for dynamic use) — for static export we use the JSON
        charts = ChartBuilder.build_all(pack, proxy_results)

        # Parse interpretation
        per_sample: list[dict] = []
        overall_narrative = ""
        citations: list[dict] = []
        if interpretation:
            try:
                per_sample = json.loads(interpretation.per_sample_json)
            except (json.JSONDecodeError, TypeError):
                per_sample = []
            overall_narrative = interpretation.overall_narrative or ""
            try:
                citations_raw = json.loads(interpretation.citations_json)
                citations = [
                    {
                        "authors": c.get("authors", ""),
                        "year": c.get("year"),
                        "title": c.get("title", ""),
                        "venue": c.get("venue", ""),
                        "doi_or_url": c.get("doi_or_url", ""),
                        "verified": c.get("verified", False),
                    }
                    for c in citations_raw
                ]
            except (json.JSONDecodeError, TypeError):
                citations = []

        # Compute proxy stats for DeepSeek
        proxy_stats = []
        for pid in proxy_ids:
            values = [r.value for r in proxy_results if r.proxy_id == pid and r.value is not None]
            if values:
                proxy_stats.append({
                    "label": _proxy_label(pack, pid),
                    "formula": next((p.formula for p in pack.proxies if p.id == pid), ""),
                    "avg": sum(values) / len(values),
                    "min": min(values),
                    "max": max(values),
                })

        return {
            "pack_name": pack.name,
            "pack_description": pack.description or "",
            "proxies": [
                {
                    "id": p.id,
                    "label": _proxy_label(pack, p.id),
                    "description": _proxy_description(pack, p.id),
                    "formula": p.formula,
                }
                for p in pack.proxies
            ],
            "sample_ids": sample_ids,
            "proxy_ids": proxy_ids,
            "proxy_labels": proxy_labels,
            "table_rows": table_rows,
            "charts": charts,
            "per_sample_interpretation": per_sample,
            "overall_narrative": overall_narrative,
            "citations": citations,
            "generation_date": datetime.now().strftime("%B %d, %Y"),
            "executive_summary": deepseek_data.get("executive_summary", "") if deepseek_data else "",
            "key_findings": deepseek_data.get("key_findings", []) if deepseek_data else [],
            "methodology_notes": deepseek_data.get("methodology_notes", "") if deepseek_data else "",
            "recommendations": deepseek_data.get("recommendations", []) if deepseek_data else [],
            "confidence_assessment": deepseek_data.get("confidence_assessment", "") if deepseek_data else "",
            "proxy_stats": proxy_stats,
        }

    @staticmethod
    def assemble_pdf(
        pack: AnalysisPack,
        analysis_rows: list[AnalysisResult],
        interpretation: InterpretationResult | None,
        user_tier: str = "free",
        deepseek_data: dict | None = None,
    ) -> bytes:
        """Render report as PDF bytes via WeasyPrint using the DeepSeek-enhanced template."""
        data = ReportAssembler.build_report_data(pack, analysis_rows, interpretation, deepseek_data)
        data["user_tier"] = user_tier

        template = _env.get_template("report_deepseek.html")
        html_str = template.render(**data)

        pdf_bytes = HTML(string=html_str).write_pdf()
        return pdf_bytes

    @staticmethod
    def assemble_docx(
        pack: AnalysisPack,
        analysis_rows: list[AnalysisResult],
        interpretation: InterpretationResult | None,
        user_tier: str = "free",
        deepseek_data: dict | None = None,
    ) -> bytes:
        """Render report as DOCX bytes via python-docx."""
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import RGBColor, Pt

        data = ReportAssembler.build_report_data(pack, analysis_rows, interpretation, deepseek_data)

        doc = Document()

        # Title
        title = doc.add_heading(f"DataProof Report — {data['pack_name']}", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Watermark for free tier
        if user_tier == "free":
            wm = doc.add_paragraph()
            wm.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = wm.add_run("Created by Mohamad Kamardin with DataProof Engine")
            run.bold = True
            run.font.color.rgb = RGBColor(0xC9, 0x7D, 0x4A)

        # Executive Summary
        if data.get("executive_summary"):
            doc.add_heading("Executive Summary", level=2)
            doc.add_paragraph(data["executive_summary"])

        # Key Findings
        if data.get("key_findings"):
            doc.add_heading("Key Findings", level=2)
            for i, finding in enumerate(data["key_findings"], 1):
                doc.add_paragraph(f"{i}. {finding}", style="List Number")

        # Methodology section
        doc.add_heading("Methodology", level=2)
        doc.add_paragraph(data["pack_description"])
        doc.add_heading("Proxies Used", level=3)
        for p in data["proxies"]:
            desc = p["description"] or "See pack documentation."
            doc.add_paragraph(f"{p['label']} ({p['formula']}): {desc}")

        # Results table
        doc.add_heading("Results", level=2)
        table = doc.add_table(rows=1, cols=1 + len(data["proxy_ids"]))
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "Sample"
        for i, pid in enumerate(data["proxy_ids"]):
            label = _proxy_label(pack, pid)
            header[i + 1].text = label

        for row_data in data["table_rows"][:20]:
            row_cells = table.add_row().cells
            row_cells[0].text = str(row_data["sample_id"])
            for i, pid in enumerate(data["proxy_ids"]):
                val = row_data.get(pid)
                unit = row_data.get(f"{pid}_unit", "")
                if val is not None:
                    row_cells[i + 1].text = f"{val:.4f} {unit}"
                else:
                    row_cells[i + 1].text = "N/A"

        # Recommendations
        if data.get("recommendations"):
            doc.add_heading("Recommendations", level=2)
            for rec in data["recommendations"]:
                doc.add_paragraph(rec, style="List Bullet")

        # Confidence
        if data.get("confidence_assessment"):
            doc.add_heading("Confidence Assessment", level=2)
            doc.add_paragraph(data["confidence_assessment"])

        # Interpretation
        if data["overall_narrative"]:
            doc.add_heading("Interpretation", level=2)
            doc.add_paragraph(data["overall_narrative"])
            if data["per_sample_interpretation"]:
                doc.add_heading("Per-Sample Classifications", level=3)
                for ps in data["per_sample_interpretation"][:10]:
                    text = f"{ps['sample_id']}: {ps['classification']} — {ps['rationale']}"
                    doc.add_paragraph(text)

        # References
        if data["citations"]:
            doc.add_heading("References", level=2)
            for c in data["citations"]:
                verified_tag = " [VERIFIED]" if c.get("verified") else " [UNVERIFIED]"
                text = (
                    f"{c.get('authors', '')} ({c.get('year', 'n.d.')}). "
                    f"{c.get('title', '')}. {c.get('venue', '')}.{verified_tag}"
                )
                if c.get("doi_or_url"):
                    text += f" {c['doi_or_url']}"
                doc.add_paragraph(text, style="List Bullet")

        # Watermark footer for free tier
        if user_tier == "free":
            for section in doc.sections:
                footer = section.footer
                fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = fp.add_run("Created by Mohamad Kamardin with DataProof Engine")
                run.bold = True
                run.font.color.rgb = RGBColor(0xC9, 0x7D, 0x4A)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()
